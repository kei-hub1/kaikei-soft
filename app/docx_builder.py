"""python-docx による Word 文書の組み立て（A4縦・横書きのビジネス文書レイアウト）。

レイアウト（上から順）:
  右寄せ  作成日（和暦）
  左寄せ  宛先（郵便番号・住所・会社名（部署名）・氏名＋敬称）
  右寄せ  差出人（登録された項目のみ）
  中央    件名（太字・12pt）
  左寄せ  頭語（拝啓）＋本文
  右寄せ  結語（敬具）
  中央    「記」                ┐
  左寄せ  書類一覧（番号・書類名・部数） │ 記書きありの場合のみ
  右寄せ  「以上」              ┘
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .models import DocumentRequest
from .placeholders import PlaceholderContext, render
from .wareki import to_wareki

FONT_NAME = "ＭＳ 明朝"
BODY_SIZE = Pt(10.5)
SUBJECT_SIZE = Pt(12)
PAGE_MARGIN = Mm(25)
LINE_SPACING = 1.3  # 本文の行間（倍率）

_THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


def build_document(req: DocumentRequest) -> Document:
    """文書作成依頼から Word 文書オブジェクトを組み立てる（保存はしない）。"""
    ctx = PlaceholderContext(customer=req.customer, sender=req.sender, created_on=req.created_on)
    subject = render(req.subject, ctx).strip()
    body = render(req.body, ctx)

    doc = Document()
    _setup_page(doc)
    _setup_normal_style(doc)
    doc.core_properties.title = subject
    doc.core_properties.author = req.sender.office.strip() or ""

    # 作成日
    _para(doc, to_wareki(req.created_on), align=WD_ALIGN_PARAGRAPH.RIGHT)
    _blank(doc)

    # 宛先
    for line in req.customer.address_lines():
        _para(doc, line)
    _blank(doc)

    # 差出人
    for line in req.sender.lines():
        _para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _blank(doc)

    # 件名
    _para(doc, subject, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=SUBJECT_SIZE)
    _blank(doc)

    # 本文（頭語＋本文、各段落は1字下げ）
    _body_paragraphs(doc, req.opening.strip(), body)

    # 結語
    if req.closing.strip():
        _para(doc, req.closing.strip(), align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 記書き
    enclosures = req.active_enclosures()
    if req.has_notes:
        _blank(doc)
        _para(doc, "記", align=WD_ALIGN_PARAGRAPH.CENTER)
        _blank(doc)
        if enclosures:
            _enclosure_table(doc, enclosures)
            _blank(doc)
        _para(doc, "以上", align=WD_ALIGN_PARAGRAPH.RIGHT)

    return doc


def save_document(req: DocumentRequest, path: Path) -> Path:
    """文書を組み立てて .docx として保存する。"""
    doc = build_document(req)
    doc.save(str(path))
    return path


# ---- 内部ヘルパー -----------------------------------------------------------


def _setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = PAGE_MARGIN
    section.top_margin = section.bottom_margin = PAGE_MARGIN


def _setup_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.size = BODY_SIZE
    _apply_font(style.element.get_or_add_rPr())
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = LINE_SPACING
    # 東アジア言語を日本語に指定（Word の校正・禁則処理のため）
    lang = style.element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        style.element.get_or_add_rPr().append(lang)
    lang.set(qn("w:eastAsia"), "ja-JP")


def _apply_font(rpr) -> None:
    """rPr 要素に MS 明朝を設定する（英数字・日本語の両方）。"""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)
    for attr in _THEME_ATTRS:
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _run(paragraph, text: str, bold: bool = False, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = size or BODY_SIZE
    _apply_font(run._element.get_or_add_rPr())
    return run


def _para(doc: Document, text: str, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=None,
          first_line_indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    if first_line_indent is not None:
        p.paragraph_format.first_line_indent = first_line_indent
    _run(p, text, bold=bold, size=size)
    return p


def _blank(doc: Document):
    p = doc.add_paragraph()
    _run(p, "")
    return p


def _body_paragraphs(doc: Document, opening: str, body: str) -> None:
    """頭語と本文を段落に分けて出力する。

    最初の段落は「拝啓　本文…」のように頭語に続けて書き、
    2段落目以降は1字下げにする。本文中の空行はそのまま空行として残す。
    """
    lines = [ln.rstrip() for ln in body.replace("\r\n", "\n").strip("\n").split("\n")]
    if not lines or all(not ln for ln in lines):
        lines = []

    first_done = False
    for line in lines:
        if not line.strip():
            _blank(doc)
            continue
        if not first_done and opening:
            _para(doc, f"{opening}　{line.strip()}")
            first_done = True
        else:
            _para(doc, line.strip(), first_line_indent=BODY_SIZE)
            first_done = True
    if not first_done and opening:
        _para(doc, opening)


def _enclosure_table(doc: Document, enclosures) -> None:
    """番号・書類名・部数の3列を罫線なしの表で並べる。"""
    widths = (Mm(12), Mm(95), Mm(25))
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_indent(table, Mm(15))
    _set_fixed_layout(table)
    for column, width in zip(table.columns, widths):
        column.width = width
    for i, enc in enumerate(enclosures, start=1):
        row = table.add_row()
        cells = row.cells
        values = (f"{i}.", enc.title.strip(), enc.quantity_label)
        for cell, width, value in zip(cells, widths, values):
            cell.width = width
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _run(p, value)
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_fixed_layout(table) -> None:
    """列幅を固定にする（Word・LibreOffice の双方で列幅を尊重させる）。"""
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_indent(table, indent) -> None:
    tbl_pr = table._tbl.tblPr
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), str(int(indent / 635)))  # EMU → twips
    ind.set(qn("w:type"), "dxa")

