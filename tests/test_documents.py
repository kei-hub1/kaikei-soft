import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app import docx_builder, service
from app.models import Customer, DocumentRequest, Enclosure, Sender
from app.pdf_converter import PdfConversionError


def _request(**overrides) -> DocumentRequest:
    base = dict(
        customer=Customer(company="株式会社サンプル商事", department="経理部", name="鈴木 一郎", honorific="様",
                          postal="530-0001", address1="大阪府大阪市北区梅田1-2-3", address2="梅田タワー12階"),
        sender=Sender(office="吉川和章税理士事務所", accountant="吉川 和章", tel="03-1234-5678"),
        created_on=dt.date(2026, 9, 17),
        subject="資料送付のご案内",
        body="平素は格別のご高配を賜り、厚く御礼申し上げます。\nさて、{宛名}宛てに下記の資料をお送りいたします。",
        has_notes=True,
        enclosures=[Enclosure("決算報告書", "1"), Enclosure("", ""), Enclosure("総勘定元帳", "一式")],
    )
    base.update(overrides)
    return DocumentRequest(**base)


def _texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


class DocxBuilderTest(unittest.TestCase):
    def test_layout_with_notes(self):
        doc = docx_builder.build_document(_request())
        texts = _texts(doc)
        nonblank = [t for t in texts if t]
        self.assertEqual(
            nonblank,
            [
                "令和8年9月17日",
                "〒530-0001", "大阪府大阪市北区梅田1-2-3", "梅田タワー12階", "株式会社サンプル商事　経理部", "鈴木 一郎　様",
                "吉川和章税理士事務所", "税理士　吉川 和章", "TEL：03-1234-5678",
                "資料送付のご案内",
                "拝啓　平素は格別のご高配を賜り、厚く御礼申し上げます。",
                "さて、鈴木 一郎　様宛てに下記の資料をお送りいたします。",
                "敬具", "記", "以上",
            ],
        )
        # 空欄項目（FAX等）が空行として残っていない：本文中で空行が連続しない
        # （段落と表を文書内の順番どおりに並べて確認する）
        sequence = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                sequence.append("".join(t.text or "" for t in child.iter(qn("w:t"))))
            elif child.tag == qn("w:tbl"):
                sequence.append("<table>")
        for a, b in zip(sequence, sequence[1:]):
            self.assertFalse(a == "" and b == "", f"空行が連続しています: {sequence}")

        table = doc.tables[0]
        rows = [[c.text for c in r.cells] for r in table.rows]
        self.assertEqual(rows, [["1.", "決算報告書", "1部"], ["2.", "総勘定元帳", "一式"]])

    def test_layout_without_notes(self):
        doc = docx_builder.build_document(_request(has_notes=False, body="本文のみ"))
        texts = _texts(doc)
        self.assertNotIn("記", texts)
        self.assertNotIn("以上", texts)
        self.assertEqual(len(doc.tables), 0)
        self.assertEqual(texts[-1], "敬具")

    def test_alignment_and_fonts(self):
        doc = docx_builder.build_document(_request())
        by_text = {p.text: p for p in doc.paragraphs if p.text}
        self.assertEqual(by_text["令和8年9月17日"].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(by_text["吉川和章税理士事務所"].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(by_text["資料送付のご案内"].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(by_text["敬具"].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(by_text["記"].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(by_text["以上"].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertIn(by_text["〒530-0001"].alignment, (None, WD_ALIGN_PARAGRAPH.LEFT))

        subject_run = by_text["資料送付のご案内"].runs[0]
        self.assertTrue(subject_run.bold)
        self.assertEqual(subject_run.font.size.pt, 12)
        rfonts = subject_run._element.rPr.find(qn("w:rFonts"))
        self.assertEqual(rfonts.get(qn("w:eastAsia")), "ＭＳ 明朝")

        normal = doc.styles["Normal"]
        self.assertEqual(normal.font.size.pt, 10.5)
        self.assertEqual(normal.element.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "ＭＳ 明朝")

    def test_page_setup(self):
        doc = docx_builder.build_document(_request())
        s = doc.sections[0]
        self.assertAlmostEqual(s.page_width.mm, 210, places=1)
        self.assertAlmostEqual(s.page_height.mm, 297, places=1)
        for m in (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin):
            self.assertAlmostEqual(m.mm, 25, places=1)

    def test_custom_opening_closing(self):
        doc = docx_builder.build_document(_request(opening="", closing="", has_notes=False, body="本文"))
        self.assertEqual([t for t in _texts(doc) if t][-1], "本文")

    def test_subject_placeholder_rendered(self):
        doc = docx_builder.build_document(_request(subject="{作成日}のご案内", has_notes=False))
        self.assertIn("令和8年9月17日のご案内", _texts(doc))


class ServiceTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.out = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_validation_messages(self):
        with self.assertRaises(service.ValidationError) as cm:
            service.validate(_request(subject=" ", body="", customer=Customer()))
        msg = str(cm.exception)
        self.assertIn("宛先", msg)
        self.assertIn("件名", msg)
        self.assertIn("本文", msg)
        with self.assertRaises(service.ValidationError):
            service.validate(_request(has_notes=True, enclosures=[]))

    def test_create_document_with_pdf(self):
        def fake_convert(docx_path, pdf_path):
            Path(pdf_path).write_bytes(b"%PDF-")
            return "Fake"

        with mock.patch.object(service, "convert_docx_to_pdf", fake_convert):
            result = service.create_document(_request(), self.out)
        self.assertTrue(result.pdf_ok)
        self.assertEqual(result.docx_path.name, "20260917_株式会社サンプル商事_資料送付のご案内.docx")
        self.assertEqual(result.pdf_path.name, "20260917_株式会社サンプル商事_資料送付のご案内.pdf")
        self.assertTrue(result.docx_path.exists())
        Document(str(result.docx_path))  # 開けること

    def test_pdf_failure_keeps_docx(self):
        with mock.patch.object(service, "convert_docx_to_pdf", side_effect=PdfConversionError("だめ")):
            result = service.create_document(_request(), self.out)
        self.assertFalse(result.pdf_ok)
        self.assertEqual(result.pdf_error, "だめ")
        self.assertTrue(result.docx_path.exists())
        self.assertIsNone(result.pdf_path)

    def test_no_overwrite_sequence_numbers(self):
        with mock.patch.object(service, "convert_docx_to_pdf", side_effect=PdfConversionError("x")):
            r1 = service.create_document(_request(), self.out)
            r2 = service.create_document(_request(), self.out)
        self.assertEqual(r2.docx_path.name, r1.docx_path.stem + "_2.docx")

    def test_filename_uses_rendered_subject_and_safe_chars(self):
        with mock.patch.object(service, "convert_docx_to_pdf", side_effect=PdfConversionError("x")):
            r = service.create_document(_request(subject="{作成日}／資料:送付?"), self.out)
        self.assertEqual(r.docx_path.name, "20260917_株式会社サンプル商事_令和8年9月17日／資料：送付？.docx")


if __name__ == "__main__":
    unittest.main()
